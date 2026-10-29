"""Deterministic customer views. Creator calibration never enters default delivery."""

import argparse
import base64
import hashlib
import html
import importlib.util
import json
from pathlib import Path
import re
import sys
import tempfile
import textwrap

HERE = Path(__file__).resolve().parent
SPEC = importlib.util.spec_from_file_location("design", HERE / "design.py")
D = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(D)

CSS = """
:root{--navy:#17324d;--ink:#182d40;--page:#f4f6f8;--paper:#fff;--muted:#40576b;--link:#164d82;--line:#cbd4dd;--focus:#b7571b;--tint:#e5edf4;--input-line:#607286;--code:#f3f5f7;--display:2.4rem;--heading:1.65rem;--subheading:1.2rem;--small:.88rem;--mobile-display:1.9rem;font-family:Aptos,Segoe UI,Arial,sans-serif;color:var(--ink);background:var(--page);font-size:17px}
*{box-sizing:border-box}body{margin:0;overflow-wrap:anywhere}header,main,footer{max-width:1080px;margin:auto;padding:28px;min-width:0}
header{padding-top:36px}h1{font-size:var(--display);line-height:1.12;max-width:25ch;letter-spacing:-.025em}
h2{margin:2rem 0 1rem;font-size:var(--heading)}h3{font-size:var(--subheading);margin:0 0 .7rem}
p,li{line-height:1.6;max-width:72ch}a{color:var(--link);text-underline-offset:.2em}
nav{display:flex;flex-wrap:wrap;gap:8px;margin-top:24px}button,input,select{font:inherit;min-height:44px;padding:10px 14px}
button{border:1px solid var(--navy);background:var(--paper);color:var(--navy);cursor:pointer;border-radius:3px;max-width:100%}
button[aria-selected=true]{color:var(--paper);background:var(--navy)}button:hover{background:var(--tint);color:var(--navy)}
:focus-visible{outline:3px solid var(--focus);outline-offset:3px}.filters{display:flex;flex-wrap:wrap;gap:16px;align-items:end}
label{display:grid;gap:6px;min-width:0;max-width:100%}input,select{border:1px solid var(--input-line);background:var(--paper);width:100%;min-width:0;max-width:100%}
article,.family{background:var(--paper);border:1px solid var(--line);padding:24px;margin:20px 0;border-radius:3px;min-width:0}
.test-input{margin:20px 0;padding:18px;background:var(--code)}.test-input pre{padding:0;background:transparent}
.checks{padding-left:24px}.checks>li{padding:10px 0}.checks details{margin:0}.checks summary{font-size:var(--small);min-height:44px}
.test-purpose{font-size:var(--subheading);line-height:1.55}.start-actions{margin:20px 0;display:flex;gap:14px;flex-wrap:wrap}
.start-actions a{display:inline-block;padding:12px 16px;border:1px solid var(--navy);border-radius:3px;text-decoration:none}
.start-actions a:first-child{background:var(--navy);color:var(--paper)}
.meta{font-size:var(--small);color:var(--muted)}.status{padding:12px 16px;background:var(--tint);line-height:1.5}
details{margin:14px 0}summary{cursor:pointer;color:var(--link);padding:10px 0;min-height:44px}
pre{font:inherit;line-height:1.55;white-space:pre-wrap;overflow-wrap:anywhere;background:var(--code);padding:16px}
.table-wrap{overflow:auto}table{border-collapse:collapse;width:100%;text-align:left;background:var(--paper)}
th,td{padding:12px;border-bottom:1px solid var(--line);vertical-align:top;line-height:1.5}
th{background:var(--navy);color:var(--paper)}dl{display:grid;grid-template-columns:minmax(0,11rem) minmax(0,1fr);gap:10px 18px}
dt{font-weight:bold}dd{margin:0;line-height:1.5}.criterion{margin:16px 0;padding-top:12px;border-top:1px solid var(--line)}
[hidden]{display:none!important}.empty{padding:24px;background:var(--paper)}footer{font-size:var(--small);color:var(--muted)}
.actions{display:flex;gap:12px;flex-wrap:wrap;margin:16px 0}.print-note{display:none}.priority-fields{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,15rem),1fr));gap:16px}
@media(max-width:650px){header,main,footer{padding:20px}h1{font-size:var(--mobile-display)}article,.family{padding:16px}dl{display:block}dd{margin:6px 0 18px}.filters{display:block}label{margin-bottom:14px}}
@media print{nav,.filters,.actions,#profile-form{display:none}section[hidden]{display:block!important}article[hidden]{display:block!important}article{break-inside:auto}.print-note{display:block}#count,#empty{display:none!important}.table-wrap{overflow:visible}}
"""
JS = """
const tabs=[...document.querySelectorAll('[role=tab]')];
function activate(tab){tabs.forEach(t=>{const on=t===tab;t.setAttribute('aria-selected',on);t.tabIndex=on?0:-1;
document.getElementById(t.getAttribute('aria-controls')).hidden=!on;});}
tabs.forEach((t,i)=>{t.addEventListener('click',()=>activate(t));t.addEventListener('keydown',e=>{
if(e.key==='ArrowRight'||e.key==='ArrowLeft'){e.preventDefault();const n=tabs[(i+(e.key==='ArrowRight'?1:-1)+tabs.length)%tabs.length];activate(n);n.focus();}});});
const search=document.getElementById('search'),dimension=document.getElementById('dimension'),track=document.getElementById('track');
function filter(){let n=0;document.querySelectorAll('article[data-search]').forEach(a=>{
const show=a.dataset.search.includes(search.value.trim().toLowerCase())&&(!dimension.value||a.dataset.dimensions.split(' ').includes(dimension.value))&&(!track.value||a.dataset.track===track.value);
a.hidden=!show;if(show)n++;});document.getElementById('count').textContent=n+' cases shown';document.getElementById('empty').hidden=n>0;}
search.addEventListener('input',filter);dimension.addEventListener('change',filter);track.addEventListener('change',filter);filter();
function route(){const id=location.hash.slice(1),target=document.getElementById(id);
if(target){const panel=target.closest('[role=tabpanel]');if(panel){activate(tabs.find(t=>t.getAttribute('aria-controls')===panel.id));
if(target.matches('article')){search.value='';dimension.value='';track.value='';filter();}requestAnimationFrame(()=>target.scrollIntoView({block:'start'}));return;}}
activate(tabs[0]);}
window.addEventListener('hashchange',route);route();
document.querySelectorAll('a[href^="#"]').forEach(a=>a.addEventListener('click',()=>{if(location.hash===a.getAttribute('href'))route();}));
document.getElementById('clear-filters').addEventListener('click',()=>{search.value='';dimension.value='';track.value='';filter();search.focus();});
document.getElementById('print-all').addEventListener('click',()=>window.print());
let expandedForPrint=[];
window.addEventListener('beforeprint',()=>{expandedForPrint=[...document.querySelectorAll('details:not([open])')];expandedForPrint.forEach(d=>d.open=true);});
window.addEventListener('afterprint',()=>{expandedForPrint.forEach(d=>d.open=false);expandedForPrint=[];});
const form=document.getElementById('profile-form');
form.addEventListener('submit',e=>{e.preventDefault();const out={},fields=[...form.querySelectorAll('[data-root]')],message=document.getElementById('profile-status');
for(const f of fields){if(!f.validity.valid){message.textContent='Enter finite nonnegative numbers; leave unused dimensions blank.';f.focus();return;}
const v=f.value.trim()===''?0:Number(f.value);if(!Number.isFinite(v)||v<0){message.textContent='Priorities must be finite nonnegative numbers.';f.focus();return;}out[f.dataset.root]=v;}
const total=Object.values(out).reduce((s,v)=>s+v,0);if(!Number.isFinite(total)||total<=0){message.textContent='Enter at least one positive priority. All-zero priorities do not define a profile.';return;}
const normalized=Object.fromEntries(Object.entries(out).map(([k,v])=>[k,v/total]));
const profile={document_type:'evaluation_customer_profile',pack_hash:form.dataset.pack,overrides:out,normalized,meaning:'Customer importance preferences only; no target score or source change'};
const url=URL.createObjectURL(new Blob([JSON.stringify(profile,null,2)+'\\n'],{type:'application/json'}));const a=document.createElement('a');a.href=url;a.download='customer_priorities.json';a.click();setTimeout(()=>URL.revokeObjectURL(url),1000);
message.textContent='Profile exported. The evaluation cases and expected behavior remain unchanged.';});
"""

COLLECTION_CSS = """
header{padding-bottom:12px}main{padding-top:8px}
.collection-summary{font-size:var(--subheading);font-weight:600;line-height:1.6;max-width:65ch;margin-top:0}
.category-group{border-top:1px solid var(--line);padding-top:28px;margin-top:32px}
.category-group h2{margin-top:0}.theme-group{padding:8px 0;border-bottom:1px solid var(--line)}
.theme-group>summary{font-weight:600}.scenario-group{margin:18px 0;padding:0 0 18px}
.scenario-group:last-child{padding-bottom:0}.scenario-group h4{font-size:1.05rem;margin:12px 0}
.principle-list{padding-left:22px}.collection-counts td:nth-child(2),.collection-counts td:nth-child(3){white-space:nowrap}
.collection-counts th:nth-child(1){width:23%}.collection-counts th:nth-child(2){width:9%;white-space:nowrap}
.collection-counts th:nth-child(3){width:16%}
.ontology-path{font-weight:600}.gap-note{background:var(--tint);padding:16px}
@media(max-width:650px){
.collection-counts,.collection-counts tbody{display:block}.collection-counts thead{position:absolute;width:1px;height:1px;overflow:hidden;clip-path:inset(50%)}
.collection-counts tr{display:grid;grid-template-columns:1fr 1fr;border-bottom:1px solid var(--line);padding:12px 0}
.collection-counts td{padding:8px 12px;border:0;font-size:.95rem}.collection-counts td:first-child,.collection-counts td:last-child{grid-column:1/-1}
.collection-counts td:first-child{font-weight:700;font-size:1.05rem}.collection-counts td[data-label]:before{content:attr(data-label) ': ';font-weight:600}
}
"""

COLLECTION_JS = """
const scenarioSelect=document.getElementById('scenario');
const originalFilter=filter;
filter=function(){originalFilter();let n=0;document.querySelectorAll('article[data-search]').forEach(a=>{
if(scenarioSelect.value&&a.dataset.family!==scenarioSelect.value)a.hidden=true;if(!a.hidden)n++;});
document.getElementById('count').textContent=n+' cases shown';document.getElementById('empty').hidden=n>0;};
[search,dimension,track].forEach(x=>x.addEventListener(x===search?'input':'change',filter));
scenarioSelect.addEventListener('change',filter);
document.getElementById('clear-filters').addEventListener('click',()=>{scenarioSelect.value='';filter();});
document.querySelectorAll('[data-open-family]').forEach(a=>a.addEventListener('click',e=>{
e.preventDefault();search.value='';dimension.value='';track.value='';scenarioSelect.value=a.dataset.openFamily;
activate(document.getElementById('tab-cases'));filter();location.hash='cases';document.getElementById('cases').scrollIntoView();}));
function revealCollectionRoute(){const target=document.getElementById(location.hash.slice(1));if(!target)return;
if(target.matches('article')){scenarioSelect.value='';filter();}
let p=target.parentElement;while(p){if(p.matches('details'))p.open=true;p=p.parentElement;}}
window.addEventListener('hashchange',revealCollectionRoute);revealCollectionRoute();filter();
document.querySelectorAll('a[href^="#"]').forEach(a=>a.addEventListener('click',()=>{
if(location.hash===a.getAttribute('href'))revealCollectionRoute();}));
"""


def collection_anchor(kind, key):
    return kind + "-" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:16]


def collection_panels(model, *, standalone=False):
    c = model["collection"]
    count = c["counts"]
    principles = D.index(c["principles"])
    dims = D.index(model["dimensions"])
    scenarios = D.index(c["scenarios"])
    def principle_link(key):
        return f'<a href="#{collection_anchor("principle", key)}">{esc(principles[key]["title"])}</a>'

    rows = "".join(
        f'<tr><td><a href="#{collection_anchor("category", cat["id"])}">{esc(cat["title"])}</a></td>'
        f'<td data-label="Tests">{cat["test_count"]}</td><td data-label="Scenario groups">{cat["scenario_count"]}</td><td>{esc(cat["rationale"])}</td></tr>'
        for cat in c["categories"])
    gaps = "".join(f'<li><strong>{esc(g["dimension"])}:</strong> {esc(g["reason"])}</li>' for g in c["gaps"])
    overview = (
        f'<p class="collection-summary">{esc(D.O.summary(c))}</p>'
        f'<p>{esc(c["purpose"])}</p>'
        f'<p>{count["categories"]} categories and {count["themes"]} themes are declared; '
        f'{count["populated_themes"]} themes have tests. '
        f'{count["exposed_requirements"]} of {count["requirements"]} requirements have designed exposure; '
        f'{count["gap_requirements"]} remain gaps. These are definitions, not executed tests or scores.</p>'
        '<h2>Categories and why they were selected</h2>'
        '<div class="table-wrap"><table class="collection-counts"><thead><tr><th>Category</th>'
        '<th>Tests</th><th>Scenario groups</th><th>Why it belongs</th></tr></thead>'
        f'<tbody>{rows}</tbody></table></div>'
        '<p class="meta">Tests use one primary category, so these counts sum to the collection total. '
        'Secondary theme/category links overlap and are not additive. Case count and priority are not performance.</p>'
        '<div class="start-actions"><a href="#landscape">Explore categories and scenarios</a>'
        '<a href="#coverage">Principles and dimensions</a></div>'
        '<h2>The design logic</h2><p class="ontology-path">Principles and target dimensions explain the categories. '
        'Themes organize scenario groups. Cases and their checks sit inside those groups.</p>'
        '<p>Selection earns its place through a meaningful distinction. Balanced controls, traceable authority '
        'and acceptable alternatives matter alongside coverage. Open a category for its specific principle links.</p>'
        f'<h2>Visible coverage gaps</h2><ul>{gaps or "<li>No requirement gaps are declared in this pack.</li>"}</ul>'
        f'<details><summary>Mapping provenance and scope limits</summary><p>{esc(c["provenance"]["note"])}</p>'
        '<p>Principle links explain the authored design rationale; they are not independent proof of principle compliance.</p>'
        f'<ul>{"".join("<li>"+esc(x)+"</li>" for x in model["limitations"])}</ul></details>'
        + ("" if standalone else
           '<p><a href="EVALUATION_DESIGN.xlsx">Editable workbook</a> · '
           '<a href="EXECUTIVE_DESIGN_BRIEF.docx">Executive overview</a></p>'))
    groups = []
    for cat in c["categories"]:
        links = "".join(f'<li><strong>{principle_link(link["principle_id"])}:</strong> '
                        f'{esc(link["explanation"])}</li>' for link in cat["principle_links"])
        themes_html = []
        for theme in c["themes"]:
            if theme["category_id"] != cat["id"]:
                continue
            group_html = []
            for sid in theme["scenario_ids"]:
                scenario = scenarios[sid]
                group_html.append(
                    f'<section class="scenario-group" id="{collection_anchor("scenario", sid)}">'
                    f'<h4>{esc(scenario["title"])}</h4><p>{esc(scenario["rationale"])}</p>'
                    f'<p>{scenario["test_count"]} tests · {scenario["check_count"]} checks</p>'
                    f'<p class="meta">Principles: {", ".join(principle_link(p) for p in scenario["principle_ids"])}</p>'
                    f'<a href="#cases" data-open-family="{esc(scenario["family_id"])}">Inspect these {scenario["test_count"]} tests</a></section>')
            themes_html.append(
                f'<details class="theme-group" id="{collection_anchor("theme", theme["id"])}">'
                f'<summary>{esc(theme["title"])} — {theme["scenario_count"]} scenario groups · {theme["test_count"]} tests</summary>'
                f'<p>{esc(theme["meaning"])}</p><p><strong>Why this theme:</strong> {esc(theme["rationale"])}</p>'
                f'<p class="meta">Target dimensions: {", ".join(esc(dims[d]["label"]) for d in theme["dimension_ids"])}</p>'
                f'{"".join(group_html) or "<p class=gap-note>No tests yet. This declared theme remains a coverage gap.</p>"}'
                '</details>')
        category_gaps = [g for g in c["gaps"] if g["requirement_id"] in cat["gap_requirement_ids"]]
        groups.append(
            f'<section class="category-group" id="{collection_anchor("category", cat["id"])}">'
            f'<h2>{esc(cat["title"])}</h2><p>{esc(cat["meaning"])}</p>'
            f'<p><strong>Why selected:</strong> {esc(cat["rationale"])}</p>'
            f'<p>{cat["test_count"]} primary tests · {cat["scenario_count"]} scenario groups</p>'
            f'<p class="meta">Target dimensions: {", ".join(esc(dims[d]["label"]) for d in cat["dimension_ids"])}</p>'
            f'<h3>Principles behind the selection</h3><ul class="principle-list">{links}</ul>'
            f'<h3>Themes and top-level scenarios</h3>{"".join(themes_html)}'
            + (f'<p class="gap-note"><strong>Uncovered:</strong> {"; ".join(esc(g["dimension"]) for g in category_gaps)}</p>'
               if category_gaps else "") + '</section>')
    principles_html = "".join(
        f'<section id="{collection_anchor("principle", p["id"])}"><h3>{esc(p["title"])}</h3>'
        f'<p>{esc(p["meaning"])}</p><p><strong>How it appears:</strong> {esc(p["expression"])}</p>'
        f'<p class="meta">{esc(p["limit"])}</p></section>' for p in c["principles"])
    def tree(parent):
        children = [d for d in model["dimensions"] if d["parent_id"] == parent]
        return "<ul>" + "".join(f'<li><strong>{esc(d["label"])}:</strong> {esc(d["definition"])}{tree(d["id"])}</li>'
                               for d in children) + "</ul>" if children else ""
    generator = "".join(f'<li><strong>{esc(d["title"])}:</strong> {esc(d["meaning"])}</li>' for d in c["generator_dimensions"])
    return {
        "overview": overview,
        "landscape": '<h2>Categories, themes and scenario groups</h2>'
                     '<p>Start with the reason for each category. Expand a theme to see its scenario groups; inspect individual tests only when needed.</p>'
                     + "".join(groups),
        "principles": '<h2>Principles behind the evaluation design</h2>'
                      '<p>These are the existing ten design principles. Category mappings are explicit authored or model-proposed rationale, not a compliance score.</p>'
                      + principles_html + '<h2>Target dimensions and subdimensions</h2>'
                      '<p>These describe the target system. They anchor the category and theme coverage above.</p>' + tree(None) +
                      '<details><summary>The separate 11 generator-quality responsibilities</summary>'
                      '<p>These evaluate the creation process and collection, not eleven scores to impose on every target.</p>'
                      f'<ol>{generator}</ol></details>'
    }


def esc(value):
    return html.escape(str(value), quote=True)


def md(value):
    return str(value).replace("&", "&amp;").replace("\\", "\\\\").replace("<", "&lt;").replace(">", "&gt;").replace(
        "[", "\\[").replace("]", "\\]").replace("|", "\\|").replace("`", "\\`").replace("#", "\\#")


def stem(value):
    return "case-" + hashlib.sha256(value.encode("utf-8")).hexdigest()[:20]


def support(ref):
    part = ref.get("excerpt")
    if part:
        return part["label"], part["text"], part["explanation"]
    path = ref["pointer"]
    label = f"Task rule {int(path.rsplit('/', 1)[1]) + 1}" if path.startswith("/context/policy/") else "Supplied source"
    return label, ref["text"], "Full source reference; precise support remains subject to the declared review status."


def searchable_text(case, dimensions):
    readable = [case["id"], case["title"], case["family_id"], case["input"],
                [dimensions[key]["label"] for key in case["dimension_ids"]],
                [{key: c[key] for key in ("statement", "applicability", "alternatives", "consequence")}
                 for c in case["criteria"]]]
    return json.dumps(readable, ensure_ascii=False).lower()


def extra_observation_sections(case):
    for key, value in case["input"]["case"]["observations"].items():
        if key != "documents":
            text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2)
            yield key.replace("_", " ").capitalize(), text


def case_markdown(case, family=None):
    lines = [f"# {md(case['title'])}", "", f"**Case:** {md(case['id'])} | **Track:** {md(case['track'])}",
             "", md(case["input"]["case"]["request"]), "", "## Expected behavior"]
    for c in case["criteria"]:
        lines += ["", f"### {md(c['id'])}", md(c["statement"]),
                  f"**Applies when:** {md(c['applicability'])}", f"**Alternatives:** {md(c['alternatives'])}",
                  f"**Why it matters:** {md(c['consequence'])}", "", "Evidence:"]
        for ref in c["evidence"]:
            label, quote, explanation = support(ref)
            lines += [f"- **{md(label)}:** {md(quote)}", f"  {md(explanation)}"]
    lines += ["", "## Source notes"]
    for doc in case["input"]["case"]["observations"].get("documents", []):
        lines += ["", f"### {md(doc['id'])} — {md(doc['title'])}", md(doc["text"])]
    for title, text in extra_observation_sections(case):
        lines += ["", f"### {md(title)}", md(text)]
    lines += ["", "## Task rules", *[f"{i}. {md(p)}" for i, p in enumerate(case["input"]["context"]["policy"], 1)]]
    if family:
        lines += ["", "## Why this family belongs", md(family.get("selection_reason", family["mechanism"]))]
        for contrast in family["contrasts"]:
            lines += [md(f"{contrast['a']} / {contrast['b']}: {contrast['expected_relation']}"),
                      "Unchanged: " + md(contrast["invariants"]), "Limit: " + md(contrast["limit"])]
    lines += ["", "## Exact machine input", "", md(json.dumps(case["input"], ensure_ascii=False, indent=2))]
    for example in case["examples"]:
        lines += ["", "## Authored worked example", md(example["response"])]
    lines += ["", "## Quality Check", "Design material only; no target execution or score is implied."]
    return "\n".join(lines) + "\n"


def download_attributes(text, filename, media_type):
    payload = base64.b64encode(text.encode("utf-8")).decode("ascii")
    return f'href="data:{media_type};charset=utf-8;base64,{payload}" download="{esc(filename)}"'


def render_html(model, *, standalone_inputs=None):
    """Embed downloads when validated candidate inputs are supplied; otherwise use bundle links."""
    standalone = standalone_inputs is not None
    candidate_download = 'href="../run/CANDIDATE_INPUTS.jsonl" download'
    office_links = '<p><a href="EVALUATION_DESIGN.xlsx">Editable workbook</a> · <a href="EXECUTIVE_DESIGN_BRIEF.docx">Executive brief</a></p>'
    download_guidance = '<p><a href="../run/CANDIDATE_INPUTS.jsonl" download>Download test inputs (JSONL)</a> · <a href="../run/ADAPTER_GUIDE.md">How to connect them to your system</a></p>'
    profile_guidance = '<p>To export priorities edited in Excel, use the bundled <a href="../run/PROFILE_GUIDE.md">standalone profile guide</a>. Only the delivery and Python are needed; no creator pack or third-party packages.</p>'
    if standalone:
        candidate_download = download_attributes(
            "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in standalone_inputs),
            "CANDIDATE_INPUTS.jsonl", "application/x-ndjson")
        office_links = ""
        download_guidance = (
            f'<p><a {candidate_download}>Download test inputs (JSONL)</a></p>'
            '<p>Send only each record\'s <code>input</code> to the system being evaluated. '
            'Keep <code>case_id</code> to associate its response. The casebook and Markdown '
            'downloads contain grading material and must stay with the reviewer.</p>')
        profile_guidance = (
            '<p>Export the priorities above as <code>customer_priorities.json</code>. '
            'This records your design preferences, not measured performance. '
            'Office exports belong to the separate optional full-delivery workflow.</p>')
    card = model["card"]
    collection = model.get("collection")
    tabs = ([("overview", "Overview"), ("landscape", "Categories & scenarios"), ("coverage", "Principles & dimensions"),
             ("cases", "Test details"), ("use", "Download & adapt")] if collection else
            [("overview", "Start here"), ("cases", "Browse tests"), ("coverage", "What this covers"), ("use", "Download & adapt")])
    nav = "".join(f'<button role="tab" id="tab-{key}" aria-controls="{key}" aria-selected="false">{label}</button>'
                  for key, label in tabs)
    articles = []
    dims = D.index(model["dimensions"])
    families = D.index(model["families"])
    track_names = {"diagnostic": "Small diagnostic tests", "worklike": "Multi-document scenarios",
                   "representative": "Representative-work tests", "reserved": "Reserved tests"}
    track_options = "".join(f'<option value="{esc(key)}">{esc(track_names[key])}</option>'
                            for key in dict.fromkeys(c["track"] for c in model["cases"]))

    def case_links(ids):
        return ", ".join(f'<a href="#{stem(cid)}">{esc(cid)}</a>' for cid in ids)

    def dimensions_for(case):
        ids = set(case["dimension_ids"])
        for key in list(ids):
            while dims[key]["parent_id"]:
                key = dims[key]["parent_id"]
                ids.add(key)
        return ids

    common_rules = {}
    for case in model["cases"]:
        for criterion in case["criteria"]:
            if criterion["scope"] == "shared":
                common_rules.setdefault(criterion["statement"], criterion)

    for case in model["cases"]:
        rules = []
        for c in case["criteria"]:
            if c["scope"] == "shared":
                continue
            support_html = "".join(f"<p><strong>{esc(label)}</strong></p><pre>{esc(quote)}</pre><p>{esc(explanation)}</p>"
                                   for label, quote, explanation in map(support, c["evidence"]))
            rules.append(f'<li>{esc(c["statement"])}'
                         f'<details><summary>Why this is required</summary><p>{esc(c["consequence"])}</p>{support_html}'
                         f'<p><strong>Other acceptable answers:</strong> {esc(c["alternatives"])}</p>'
                         f'<p class="meta">Applies when: {esc(c["applicability"])}</p></details></li>')
        examples = "".join(f'<details><summary>Authored example: {esc(a["kind"])}</summary>'
                           f'<p>{esc(a["response"])}</p><p class="meta">Illustration, not a target result.</p></details>'
                           for a in case["examples"])
        search_text = searchable_text(case, dims)
        family = families[case["family_id"]]
        related = ", ".join(f'<a href="#{stem(cid)}">Open the related test ({esc(cid)})</a>'
                            for cid in family["case_ids"] if cid != case["id"])
        comparisons = "".join(f'<p><strong>{esc(p["a"])} / {esc(p["b"])}:</strong> {esc(p["expected_relation"])}</p>'
                              f'<p>Unchanged: {esc(p["invariants"])}</p><p>Limit: {esc(p["limit"])}</p>'
                              for p in family["contrasts"] if case["id"] in (p["a"], p["b"]))
        documents = case["input"]["case"]["observations"].get("documents", [])
        source_html = "".join(f'<h4>{esc(d["title"])} <span class="meta">({esc(d["id"])})</span></h4><pre>{esc(d["text"])}</pre>' for d in documents)
        policies = "".join(f'<li>{esc(s)}</li>' for s in case["input"]["context"]["policy"])
        task = case["input"]["case"]["request"]
        purpose = family.get("selection_reason", family["mechanism"])
        limits = "".join(f"<li>{esc(p['limit'])}</li>" for p in family["contrasts"])
        for title, text in extra_observation_sections(case):
            source_html += f'<h4>{esc(title)}</h4><pre>{esc(text)}</pre>'
        case_download = f'href="cases/{stem(case["id"])}.md"'
        if standalone:
            case_download = download_attributes(case_markdown(case, family), stem(case["id"]) + ".md",
                                                "text/markdown")
        articles.append(f'<article id="{stem(case["id"])}" data-search="{esc(search_text)}" '
                        f'data-track="{esc(case["track"])}" data-dimensions="{esc(" ".join(sorted(dimensions_for(case))))}"'
                        + (f' data-family="{esc(case["family_id"])}"' if collection else "") + '>'
                        f'<h3>{esc(case["title"])}</h3><p>{esc(purpose)}</p>'
                        f'<div class="test-input"><h4>Task and source material</h4><p>{esc(task)}</p>'
                        f'{source_html}<details><summary>Instructions used with this input</summary><ol>{policies}</ol>'
                        f'<p>{esc(case["input"]["context"]["tool_boundary"])}</p></details></div>'
                        f'<h4>What a good response must do</h4><ol class="checks">{"".join(rules)}</ol>'
                        f'<p class="meta">Apply any collection-wide rules listed on <a href="#overview">{esc(tabs[0][1])}</a>. Different wording is welcome when the meaning is preserved.</p>'
                        f'<details><summary>Compare with the related test</summary>'
                        f'{comparisons}<p>Related case: {related}</p></details>'
                        f'<details><summary>Test limits and technical details</summary><ul>{limits}</ul>'
                        f'<p>Reference: {esc(case["id"])}. Type: {esc(case["track"])}. Definition status: {esc(case["status"])}.</p>'
                        f'<pre>{esc(json.dumps(case["input"],ensure_ascii=False,indent=2))}</pre></details>'
                        f'{examples}<a {case_download}>Download this test as Markdown</a></article>')
    dimensions = "".join(f'<option value="{esc(d["id"])}">{esc(d["label"])}</option>'
                         for d in model["dimensions"])
    requirements = D.index(model["requirements"])
    coverage = "".join(f'<tr><td>{esc(dims[requirements[c["requirement_id"]]["dimension_id"]]["label"])}</td>'
                       f'<td>{esc(c["status"])}</td><td>{case_links(c["case_ids"]) or "None"}</td>'
                       f'<td>{esc(c["reason"])}</td></tr>' for c in model["coverage"])
    hierarchy = "".join(f'<h3>{esc(d["label"])} <span class="meta">{esc(d["id"])}</span></h3><p>{esc(d["definition"])}</p>'
                        + "".join(f'<p><strong>{esc(child["id"])} — {esc(child["label"])}:</strong> {esc(child["definition"])}</p>'
                                  for child in model["dimensions"] if child["parent_id"] == d["id"])
                        for d in model["dimensions"] if d["parent_id"] is None)
    family_rows = "".join(f'<div class="family"><h3>{esc(f["title"])}</h3><p>{esc(f.get("selection_reason", f["mechanism"]))}</p>'
                          f'<p>{case_links(f["case_ids"])}</p>'
                          + "".join(f'<p>{esc(c["expected_relation"])}</p><p class="meta">{esc(c["limit"])}</p>' for c in f["contrasts"])
                          + '</div>' for f in model["families"])
    priority_fields = "".join(f'<label>{esc(d["label"])}<input type="number" min="0" step="any" inputmode="decimal" data-root="{esc(d["id"])}" aria-label="{esc(d["label"])} priority"></label>'
                              for d in model["dimensions"] if d["parent_id"] is None)
    limitations = "".join(f"<li>{esc(s)}</li>" for s in model["limitations"])
    exclusions = "".join(f"<li>{esc(s)}</li>" for s in card["excluded"])
    common = "".join(f'<li>{esc(c["statement"])} <details><summary>How to apply this rule</summary>'
                     f'<p>{esc(c["alternatives"])}</p><p>{esc(c["consequence"])}</p></details></li>' for c in common_rules.values())
    coverage_summary = "".join(f'<li><strong>{esc(d["label"])}:</strong> {esc(d["definition"])}</li>'
                               for d in model["dimensions"] if d["parent_id"] is None)
    content = {
        "overview": f'<h2>Evaluate the behavior you need from your AI system</h2><p class="test-purpose">{esc(card["job"])}</p>'
                    f'<p>You get {len(model["cases"])} ready-to-inspect test inputs. Each includes the source text and a checklist of what a good response must preserve.</p>'
                    f'<div class="start-actions"><a href="#cases">Browse the tests</a><a {candidate_download}>Download all test inputs</a></div>'
                    '<h2>How to use a test</h2><ol><li>Choose a test relevant to your system and read its source text.</li>'
                    '<li>Use the expected-answer checklist to define the behavior you want. If you run your system, give it only the task and source input.</li>'
                    '<li>Compare its response with the checklist. Open the evidence explanation when a judgment is unclear.</li></ol>'
                    '<p>You can use the collection to define requirements, review examples with your team, or run evaluations yourself. Running and scoring are optional.</p>'
                    f'<h2>What the tests look for</h2><ul>{coverage_summary}</ul>'
                    f'<h2>Rules across this collection</h2><ul>{common}</ul>'
                    f'<h2>About this set</h2><p>{esc(card["scope"])}</p>'
                    '<p>This draft helps examine the behaviors described here. It does not establish that your system is ready for production.</p>'
                    f'<details><summary>Scope limits and review status</summary><ul>{exclusions}</ul><ul>{limitations}</ul></details>'
                    f'{office_links}',
        "cases": f'<h2>Choose a behavior to check</h2><p>Read the task and source first, then see what the response should preserve. Related tests help compare behavior when the input changes.</p>'
                 f'<div class="filters"><label>Search tests<input id="search" type="search" placeholder="Search titles, source text or expected behavior"></label>'
                 f'<label>What do you want to check?<select id="dimension"><option value="">All behaviors</option>{dimensions}</select></label>'
                 f'<label>Test scope<select id="track"><option value="">All tests</option>{track_options}</select></label>'
                 '<button id="clear-filters" type="button">Clear filters</button></div>'
                 f'<p id="count" role="status" aria-live="polite"></p><p id="empty" class="empty" hidden>No cases match. Clear the search or choose another dimension.</p>{"".join(articles)}',
        "coverage": f'<h2>What each dimension checks</h2>{hierarchy}<h2>Designed exposure</h2>'
                    '<p>Primary probes target a distinction; secondary exposure supplies additional context. Shared rules, when present, do not add new behavior coverage.</p>'
                    f'<div class="table-wrap"><table><thead><tr><th>Dimension</th><th>Status</th><th>Cases</th><th>Basis</th></tr></thead><tbody>{coverage}</tbody></table></div>'
                    f'<h2>Why these families belong</h2>{family_rows}',
        "use": '<h2>Take the tests into your workflow</h2><p>Download the inputs to use with your own system. Keep this guide and its answer criteria for the person reviewing the responses.</p>'
               f'{download_guidance}'
               '<h3>Choose what matters most to you</h3><p>Optionally set relative importance for the behaviors below. For example, give factual accuracy more importance than brevity. These preferences do not grade a model or change the expected answers. Leave everything blank if you do not need weights.</p>'
               f'<form id="profile-form" data-pack="{esc(model.get("pack_hash",""))}" novalidate><div class="priority-fields">{priority_fields}</div>'
               '<div class="actions"><button type="submit">Export priority profile</button></div><p id="profile-status" role="status" aria-live="polite">No priority profile set.</p></form>'
               f'{profile_guidance}'
               '<p>Keep source rights and task policy explicit when adapting. Names and dates may change only when the expected behavior and evidence remain aligned.</p>'
    }
    if collection:
        panels = collection_panels(model, standalone=standalone)
        content["overview"] = panels["overview"]
        content["landscape"] = panels["landscape"]
        content["coverage"] = panels["principles"] + '<h2>Requirement coverage ledger</h2>' + (
            '<div class="table-wrap"><table><thead><tr><th>Dimension</th><th>Status</th><th>Cases</th><th>Basis</th></tr></thead>'
            f'<tbody>{coverage}</tbody></table></div>')
        content["use"] += '<div class="actions"><button id="print-all" type="button">Print full collection</button></div>'
        scenario_options = "".join(f'<option value="{esc(s["family_id"])}">{esc(s["title"])}</option>' for s in collection["scenarios"])
        content["cases"] = content["cases"].replace(
            '<label>Test scope', f'<label>Scenario group<select id="scenario"><option value="">All scenario groups</option>{scenario_options}</select></label><label>Test scope')
    sections = "".join(f'<section id="{key}" role="tabpanel" aria-labelledby="tab-{key}">{content[key]}</section>' for key, _ in tabs)
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; script-src 'unsafe-inline'; img-src data:; base-uri 'none'; form-action 'none'">
<title>{esc(card["title"])}</title><style>{CSS + (COLLECTION_CSS if collection else "")}</style></head><body>
<!-- THESIS: relevant evaluations first; supporting depth on demand. OWN-WORLD: navy, white, work-document typography.
STORY: understand scope, find a case, inspect expected behavior, optionally adapt/export.
FIRST VIEWPORT: title, honest private status, purpose and four navigation choices.
FORM: plan-specified read/operate document, extending the workbook identity; no concept tournament. -->
<header><h1>{esc(card["title"])}</h1><p class="meta">Draft evaluation collection · Test inputs and expected behavior</p>
<nav role="tablist" aria-label="Evaluation collection">{nav}</nav>{'' if collection else '<div class="actions"><button id="print-all" type="button">Print full collection</button></div>'}
<p class="print-note">Full collection print: all {len(model["cases"])} cases included, regardless of screen filters. Supporting details are expanded for printing.</p></header><main>{sections}</main>
<footer>Evaluation definitions. Sources and expectations remain separate from target inputs.</footer><script>{JS + (COLLECTION_JS if collection else "")}</script></body></html>"""


def workbook_detail_entries(case, families):
    entries = [("ID", case["id"]), ("Title", case["title"]), ("Request", case["input"]["case"]["request"]),
               ("Status", case["status"]), ("Task boundary", case["input"]["context"]["tool_boundary"])]
    entries += [(f"Rule {i}", p) for i, p in enumerate(case["input"]["context"]["policy"], 1)]
    documents = case["input"]["case"]["observations"].get("documents", [])
    entries += [(f"Source {doc['id']} — {doc['title']}", doc["text"]) for doc in documents]
    extra = {k: v for k, v in case["input"]["case"]["observations"].items() if k != "documents"}
    if extra:
        entries.append(("Other observations", json.dumps(extra, ensure_ascii=False, indent=2)))
    for c in case["criteria"]:
        for field, label in [("statement", ""), ("role", "Role"), ("scope", "Scope"), ("review", "Review"),
                             ("mode", "Judgment"), ("applicability", "Applies when"),
                             ("consequence", "Consequence"), ("alternatives", "Alternatives")]:
            entries.append((c["id"] + (" " + label if label else ""), c[field]))
        for i, ref in enumerate(c["evidence"], 1):
            label, quote, explanation = support(ref)
            entries += [(f"{c['id']} Evidence {i}", label + "\n" + quote),
                        (f"{c['id']} Support {i}", explanation)]
    f = families[case["family_id"]]
    entries += [("Family", f["title"]), ("Why included", f.get("selection_reason", f["mechanism"]))]
    for i, contrast in enumerate(f["contrasts"], 1):
        entries += [(f"Comparison {i}", f"{contrast['a']} / {contrast['b']}: {contrast['expected_relation']}"),
                    (f"Invariants {i}", contrast["invariants"]), (f"Limit {i}", contrast["limit"])]
    for example in case["examples"]:
        entries.append(("Authored example " + example["id"], example["response"]))
    return entries


def chunks(value, maximum=500):
    # Bound both characters and line count so an Excel row cannot silently clip newlines.
    remaining = value
    result = []
    while remaining:
        limit = min(maximum, len(remaining))
        newlines = [i for i, ch in enumerate(remaining[:limit]) if ch == "\n"]
        if len(newlines) >= 12:
            limit = newlines[11] + 1
        result.append(remaining[:limit])
        remaining = remaining[limit:]
    return result or [""]


def collection_sheet_rows(model, detail=False):
    entries = [("Title", model["card"]["title"]), ("Target task", model["card"]["job"]), ("Scope", model["card"]["scope"])]
    entries += D.O.overview_entries(model, detail=detail)
    return [(f"{label} {i}/{len(chunks(value))}", part)
            for label, value in entries for i, part in enumerate(chunks(value), 1)]


def collection_principle_rows(model):
    return [(principle["title"], part) for principle in model["collection"]["principles"]
            for part in chunks(principle["meaning"] + "\n" + principle["expression"] + "\n" + principle["limit"])]


def family_sheet_rows(model):
    rows = []
    for i, family in enumerate(model["families"], 1):
        entries = [("Title", family["title"]), ("Cases", ", ".join(family["case_ids"])),
                   ("Why included", family.get("selection_reason", family["mechanism"])),
                   ("What changes", "\n".join(c["expected_relation"] for c in family["contrasts"])),
                   ("Limits", "\n".join(c["limit"] for c in family["contrasts"]))]
        for label, value in entries:
            parts = chunks(value)
            rows += [(f"Family {i} / {label} {j}/{len(parts)}", part or None) for j, part in enumerate(parts, 1)]
    return rows


def coverage_sheet_rows(model):
    rows = []
    for i, coverage in enumerate(model["coverage"], 1):
        for label, value in [("ID", coverage["requirement_id"]), ("Status", coverage["status"]),
                             ("Cases", ", ".join(coverage["case_ids"])), ("Reason", coverage["reason"])]:
            parts = chunks(value)
            rows += [(f"Requirement {i} / {label} {j}/{len(parts)}", part or None) for j, part in enumerate(parts, 1)]
    return rows


def write_workbook(model, pack_hash, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    wb.remove(wb.active)
    names = ["Start here", "Dimensions", "Principles", "Priorities", "Eval catalogue",
             "Stress families", "Coverage selection", "Use adaptation"]
    for name in names:
        wb.create_sheet(name)

    def add(sheet, values):
        ws = wb[sheet]
        D.require(all(not isinstance(v, str) or len(v) <= 32767 for v in values),
                  "Excel field exceeds cell limit; split the field explicitly, never truncate")
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Untrusted source strings never become spreadsheet formulas.

    if "collection" in model:
        add("Start here", ["Topic", "Collection overview"])
        for row in collection_sheet_rows(model):
            add("Start here", list(row))
        wb.create_sheet("Collection map", 1)
        add("Collection map", ["Field", "Category, theme and scenario rationale"])
        for row in collection_sheet_rows(model, detail=True):
            add("Collection map", list(row))
    else:
        for row in [("Topic", "Definition"), ("Title", model["card"]["title"]), ("Job", model["card"]["job"]),
                    ("Scope", model["card"]["scope"]), ("Status", model["status"]),
                    ("Cases", str(len(model["cases"])))]:
            add("Start here", row)
    for number, limit in enumerate(model["limitations"], 1):
        for part, text in enumerate(chunks(limit), 1):
            add("Start here", [f"Limit {number}, part {part}", text])
    add("Dimensions", ["ID", "Parent", "Dimension", "Meaning"])
    for d in model["dimensions"]:
        add("Dimensions", [d["id"], d["parent_id"] or "", d["label"], d["definition"]])
    if "collection" in model:
        add("Principles", ["Principle", "Meaning and required expression"])
        for row in collection_principle_rows(model):
            add("Principles", list(row))
    else:
        for row in [["Principle", "Practical meaning"], ["Source support", "Expected behavior follows supplied sources and task policy."],
                    ["Valid alternatives", "Equivalent meaning can be expressed differently."],
                    ["Optional scoring", "No target performance is implied by case counts or priorities."],
                    ["Honest scope", "Designed exposure is separate from representative coverage."]]:
            add("Principles", row)
    add("Priorities", ["Dimension ID", "Dimension", "Customer priority", "Normalized importance"])
    roots = [d for d in model["dimensions"] if d["parent_id"] is None]
    end = len(roots) + 1
    for i, d in enumerate(roots, 2):
        add("Priorities", [d["id"], d["label"], None, None])
        wb["Priorities"].cell(i, 4, f'=IF(C{end + 1}>0,C{i}/C{end + 1},"")')
        wb["Priorities"].cell(i, 3).fill = PatternFill("solid", fgColor="DBEAFE")
        wb["Priorities"].cell(i, 4).number_format = "0.0%"
    add("Priorities", ["Profile status", None])
    wb["Priorities"].cell(end + 1, 3, f'=SUM(C2:C{end})')
    wb["Priorities"].cell(end + 1, 2, f'=IF(SUM(C2:C{end})>0,"Preferences only - not scores","No priority profile set")')
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
    dv.errorTitle, dv.error, dv.showErrorMessage = "Invalid priority", "Enter a nonnegative number or leave blank.", True
    wb["Priorities"].add_data_validation(dv)
    dv.add(f"C2:C{end}")
    add("Eval catalogue", ["Case", "Task", "Track", "Dimensions", "Status", "Detail sheet"])
    for num, case in enumerate(model["cases"], 1):
        detail = f"Case {num:03}"
        wb.create_sheet(detail)
        add("Eval catalogue", [case["id"], case["title"], case["track"], ", ".join(case["dimension_ids"]), case["status"], detail])
        wb["Eval catalogue"].cell(num + 1, 6).hyperlink = f"#'{detail}'!A1"
        add(detail, ["Field", "Complete content (numbered continuations when needed)"])
        entries = workbook_detail_entries(case, D.index(model["families"]))
        for label, value in entries:
            parts = chunks(value)
            for i, chunk in enumerate(parts, 1):
                add(detail, [f"{label} {i}/{len(parts)}", chunk])
    add("Stress families", ["Family / field", "Complete content (numbered continuations)"])
    for row in family_sheet_rows(model):
        add("Stress families", list(row))
    add("Coverage selection", ["Requirement / field", "Complete content (numbered continuations)"])
    for row in coverage_sheet_rows(model):
        add("Coverage selection", list(row))
    for row in [["Topic", "Instruction"], ["Read", "Casebook: CASEBOOK.html. Full inputs and source evidence are available on demand."],
                ["Run", "Only run/CANDIDATE_INPUTS.jsonl is candidate material. No target invocation included."],
                ["Edit", "Save an edited workbook copy outside the original package. Export priorities in the casebook, or use run/PROFILE_GUIDE.md and the bundled stdlib-only export_priorities.py. Other edits are local annotations."]]:
        add("Use adaptation", row)
    for number, limit in enumerate(model["limitations"], 1):
        for part, text in enumerate(chunks(limit), 1):
            add("Use adaptation", [f"Limit {number}, part {part}", text])
    meta = wb.create_sheet("_profile")
    meta.append(["pack_hash", pack_hash])
    meta.append(["model_hash", D.digest(model)])
    meta.sheet_state = "hidden"
    payload = wb.create_sheet("_inputs")
    payload.append(["Case ID", "Part", "Exact candidate JSON"])
    for case in model["cases"]:
        for i, part in enumerate(chunks(json.dumps(case["input"], ensure_ascii=False, indent=2)), 1):
            payload.append([case["id"], i, part])
            payload.cell(payload.max_row, 3).data_type = "s"
    payload.sheet_state = "hidden"
    for ws in wb:
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else "A1"
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = 28 if letter == "A" else 70
        if ws.title == "Priorities":
            ws.column_dimensions["C"].width = 22
            ws.column_dimensions["D"].width = 24
        for row in ws:
            lines = 1
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Aptos", size=11, color="17324D")
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="17324D")
                    cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
                width = ws.column_dimensions[cell.column_letter].width
                lines = max(lines, sum(max(1, len(textwrap.wrap(part, max(10, int(width * .85)))))
                                       for part in str(cell.value or "").split("\n")))
            height = max(30, 16 * lines + 10)
            D.require(height <= 409 or ws.sheet_state == "hidden",
                      f"Workbook row would clip content: {ws.title} row {row[0].row}")
            ws.row_dimensions[row[0].row].height = min(409, height)
    wb.save(path)


def write_docx(model, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.65)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Aptos", Pt(11)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("17324D")
    doc.add_heading(model["card"]["title"], 0)
    doc.add_paragraph("Private design preview | " + model["status"] + " | No target scores")
    for title, text in [("Purpose", model["card"]["job"]), ("Scope", model["card"]["scope"])]:
        doc.add_heading(title, 1)
        doc.add_paragraph(text)
    doc.add_heading("What the collection contains", 1)
    if "collection" in model:
        for label, text in D.O.overview_entries(model, detail=False):
            p = doc.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(text)
    else:
        doc.add_paragraph(f'{len(model["cases"])} cases across {len({c["family_id"] for c in model["cases"]})} families. '
                          "Each case provides input and expected behavior; evidence is available in the casebook.")
        for d in model["dimensions"]:
            if d["parent_id"] is None:
                doc.add_paragraph(d["label"] + ": " + d["definition"])
    doc.add_heading("How to use it", 1)
    doc.add_paragraph(("Start with the collection overview and category rationale, then explore scenario groups and individual cases. "
                      if "collection" in model else
                      "Browse relevant cases, inspect their expected behavior and source evidence, then adapt the scope or priorities. ") +
                      "The workbook and casebook are operator material. Candidate-only inputs are provided for optional execution.")
    doc.add_heading("Limits and evidence status", 1)
    for line in model["limitations"] + model["card"]["excluded"]:
        doc.add_paragraph(line)
    doc.add_heading("Quality Check", 1)
    doc.add_paragraph("Generated from the same customer projection as the workbook and casebook. Mechanical parity is separate from semantic acceptance.")
    doc.save(path)


def compile_delivery(pack, destination):
    model = D.project(pack)
    destination = Path(destination).resolve()
    D.require(not destination.exists(), "Destination exists; create a new versioned delivery")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".design-", dir=destination.parent) as temporary:
        work = Path(temporary)
        for folder in ("read/cases", "run", "judge"):
            (work / folder).mkdir(parents=True)
        (work / "read" / "CASEBOOK.html").write_text(render_html(model), encoding="utf-8")
        for c in model["cases"]:
            (work / "read" / "cases" / (stem(c["id"]) + ".md")).write_text(
                case_markdown(c, D.index(model["families"])[c["family_id"]]), encoding="utf-8")
        D.write_new(work / "read" / "CUSTOMER_MODEL.json", model)
        start = f"# {md(model['card']['title'])}\n\nPrivate {model['status']} preview. No target scores.\n\n{md(model['card']['job'])}\n\n"
        if "collection" in model:
            start += md(D.O.summary(model["collection"])) + "\n\n"
            start += "[Collection overview](CASEBOOK.html) · [Categories and scenarios](CASEBOOK.html#landscape) · "
            D.write_new(work / "read" / "COLLECTION_OVERVIEW.json", model["collection"])
        else:
            start += "[Browse cases](CASEBOOK.html) · "
        start += "[Workbook](EVALUATION_DESIGN.xlsx) · [Executive brief](EXECUTIVE_DESIGN_BRIEF.docx)\n\n"
        start += "## Material limitations\n\n" + "\n".join("- " + md(s) for s in model["limitations"]) + "\n"
        (work / "read" / "START_HERE.md").write_text(start, encoding="utf-8")
        candidate_rows = D.candidates(pack)
        (work / "run" / "CANDIDATE_INPUTS.jsonl").write_text(
            "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in candidate_rows), encoding="utf-8")
        D.write_new(work / "run" / "INPUT_CONTRACT.json", D.obj(case_id=D.T, input=D.INPUT))
        (work / "run" / "ADAPTER_GUIDE.md").write_text(
            "# Optional execution\n\nEach JSONL record contains case_id and input. Map context.task, context.policy and "
            "context.tool_boundary as governing instructions; pass case.request and all case.observations as task data. "
            "Keep source-embedded instructions in the data boundary. Use one fresh invocation per case. "
            "Do not send criteria, siblings, workbook, casebook or creator notes to the target.\n\n"
            "First identify the interface named by the case: a full agent invocation, a deterministic component, or a response-injection fixture. "
            "These are different execution tracks. A named but undelivered fixture adapter is a prerequisite, not an executable supplied by this package. "
            "For composite systems, map the artifact and operator context to that system's actual input/configuration interface rather than treating every record as a chat prompt.\n\n"
            "No provider endpoint, credentials or execution are bundled. Blocked and reserved cases are omitted explicitly in manifest counts. "
            "The manual-v2 direct documents field was explicitly mapped to observations.documents on import.\n",
            encoding="utf-8")
        (work / "run" / "export_priorities.py").write_bytes((HERE / "profile_export.py").read_bytes())
        (work / "run" / "PROFILE_GUIDE.md").write_text(
            "# Export your priorities\n\nUse & adapt in the casebook exports a JSON profile directly, entirely offline.\n\n"
            "For workbook edits: keep the original delivery intact, save a copy as `edited_priorities.xlsx` next to the delivery folder, "
            "and edit only the blue Customer priority cells. From the delivery root, run:\n\n"
            "```powershell\npython .\\run\\export_priorities.py --workbook ..\\edited_priorities.xlsx --output ..\\customer_priorities.json\n```\n\n"
            "Python 3.10+ standard library only; no package install, creator pack or creator checkout is required. "
            "The helper reads the bundled public customer metadata and manifest, validates pack identity and requires finite, "
            "nonnegative values with at least one positive priority. Blank cells mean zero when at least one value is set. "
            "A new output file is required. Values normalize across these root dimensions only. "
            "It neither changes cases/criteria nor computes a model-performance score. "
            "Source/content edits require a new reviewed collection. Manifest hashes detect accidental drift under trusted local use, not malicious rewriting.\n",
            encoding="utf-8")
        D.write_new(work / "judge" / "CRITERIA.json", [
            {"case_id": c["id"], "criteria": c["criteria"]} for c in model["cases"]])
        D.write_new(work / "judge" / "WORKED_EXAMPLES.json", [
            {"case_id": c["id"], "examples": c["examples"]} for c in model["cases"] if c["examples"]])
        write_workbook(model, D.digest(pack), work / "read" / "EVALUATION_DESIGN.xlsx")
        write_docx(model, work / "read" / "EXECUTIVE_DESIGN_BRIEF.docx")
        manifest = {"document_type": "evaluation_design_delivery", "version": "2.3.0",
                    "status": "PRIVATE_PREVIEW", "pack_status": pack["review"]["status"],
                    "pack_hash": D.digest(pack), "customer_model_hash": D.digest(model),
                    "cases": len(model["cases"]), "candidate_cases": len(candidate_rows),
                    "omitted_candidate_ids": [c["id"] for c in pack["cases"] if c["status"] == "blocked" or c["track"] == "reserved"],
                    "target_executed": False, "sharing_authorized": False,
                    "entrypoints": {"read": "read/CASEBOOK.html", "workbook": "read/EVALUATION_DESIGN.xlsx",
                                    "brief": "read/EXECUTIVE_DESIGN_BRIEF.docx", "inputs": "run/CANDIDATE_INPUTS.jsonl"},
                    "files": {str(p.relative_to(work)).replace("\\", "/"): hashlib.sha256(p.read_bytes()).hexdigest()
                              for p in sorted(work.rglob("*")) if p.is_file()}}
        D.write_new(work / "manifest.json", manifest)
        verify(work)
        work.rename(destination)
    return manifest


def verify(folder):
    from openpyxl import load_workbook
    from docx import Document
    folder = Path(folder).resolve()
    manifest = D.load(folder / "manifest.json")
    D.require(manifest["status"] == "PRIVATE_PREVIEW" and not manifest["sharing_authorized"], "Unsupported delivery authority")
    actual = {str(p.relative_to(folder)).replace("\\", "/") for p in folder.rglob("*") if p.is_file()}
    D.require(actual == set(manifest["files"]) | {"manifest.json"}, "Missing/unmanifested package file")
    for name, sha in manifest["files"].items():
        p = (folder / name).resolve()
        D.require(p.is_relative_to(folder) and p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest() == sha,
                  "Package hash/path mismatch")
    model = D.load(folder / "read" / "CUSTOMER_MODEL.json")
    D.require(D.digest(model) == manifest["customer_model_hash"], "Customer model drift")
    rows = [json.loads(line) for line in (folder / "run" / "CANDIDATE_INPUTS.jsonl").read_text("utf-8").splitlines()]
    expected = [{"case_id": c["id"], "input": c["input"]} for c in model["cases"]
                if c["id"] not in manifest["omitted_candidate_ids"]]
    D.require(rows == expected and len(rows) == manifest["candidate_cases"], "Candidate projection drift")
    if manifest["version"] == "2.0.0" and "families" not in model:
        return {**manifest, "verification_scope": "Historical renderer: file hashes/model binding/candidate parity only; no new semantic-parity claim"}
    D.require(manifest["version"] in ("2.1.0", "2.2.0", "2.3.0"), "Unsupported delivery version")
    D.require((folder / "read" / "CASEBOOK.html").read_text("utf-8") == render_html(model), "HTML projection drift")
    workbook = load_workbook(folder / "read" / "EVALUATION_DESIGN.xlsx", data_only=False)
    D.require(workbook["_profile"]["B1"].value == manifest["pack_hash"] and
              workbook["_profile"]["B2"].value == D.digest(model), "Workbook binding drift")
    if manifest["version"] == "2.3.0":
        D.require(list(workbook["Stress families"].iter_rows(min_row=2, values_only=True)) == family_sheet_rows(model),
                  "Workbook family content changed or clipped")
        D.require(list(workbook["Coverage selection"].iter_rows(min_row=2, values_only=True)) == coverage_sheet_rows(model),
                  "Workbook coverage content changed or clipped")
    if "collection" in model:
        D.require(D.load(folder / "read" / "COLLECTION_OVERVIEW.json") == model["collection"], "Collection overview drift")
        for sheet, detail in (("Start here", False), ("Collection map", True)):
            expected_rows = collection_sheet_rows(model, detail)
            actual_rows = list(workbook[sheet].iter_rows(min_row=2, values_only=True))
            D.require(actual_rows[:len(expected_rows)] == expected_rows, "Workbook collection rationale/counts changed")
        D.require(list(workbook["Principles"].iter_rows(min_row=2, values_only=True)) == collection_principle_rows(model),
                  "Workbook principle framework changed")
    for number, case in enumerate(model["cases"], 1):
        families = D.index(model["families"])
        D.require((folder / "read" / "cases" / (stem(case["id"]) + ".md")).read_text("utf-8") == case_markdown(case, families[case["family_id"]]),
                  "Case card projection drift")
        ws = workbook[f"Case {number:03}"]
        payload = [(part, value) for cid, part, value in workbook["_inputs"].iter_rows(min_row=2, values_only=True) if cid == case["id"]]
        D.require([i for i, _ in payload] == list(range(1, len(payload) + 1)), "Workbook input chunks out of order")
        restored = json.loads("".join(value for _, value in payload))
        D.require(restored == case["input"], "Workbook input clipped/changed")
        expected_rows = []
        for label, value in workbook_detail_entries(case, families):
            parts = chunks(value)
            expected_rows += [(f"{label} {i}/{len(parts)}", part) for i, part in enumerate(parts, 1)]
        D.require(list(ws.iter_rows(min_row=2, values_only=True)) == expected_rows,
                  f"Workbook semantic field missing or altered: {case['id']}")
    document = Document(folder / "read" / "EXECUTIVE_DESIGN_BRIEF.docx")
    paragraphs = [p.text for p in document.paragraphs]
    for content in (model["card"]["title"], model["card"]["job"], model["card"]["scope"], *model["limitations"]):
        D.require(content in paragraphs, "Executive brief content drift")
    if "collection" in model:
        for label, text in D.O.overview_entries(model, detail=False):
            D.require(label + ": " + text in paragraphs, "Executive collection rationale/counts changed")
    return manifest


def import_profile(workbook, pack):
    from openpyxl import load_workbook
    D.checked(pack)
    wb = load_workbook(workbook, data_only=False)
    D.require(wb["_profile"]["B1"].value == D.digest(pack), "Workbook belongs to a different/stale pack")
    roots = [d for d in pack["dimensions"] if d["parent_id"] is None]
    ws = wb["Priorities"]
    D.require(ws.max_row == len(roots) + 2, "Unexpected priority rows")
    overrides = {}
    for i, root in enumerate(roots, 2):
        D.require(ws.cell(i, 1).value == root["id"], "Priority dimension/order changed")
        overrides[root["id"]] = 0 if ws.cell(i, 3).value is None else ws.cell(i, 3).value
    return D.customer_profile(pack, overrides)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("action", choices=("compile", "verify", "import-profile"))
    p.add_argument("--pack", type=Path)
    p.add_argument("--output", type=Path)
    p.add_argument("--package", type=Path)
    p.add_argument("--workbook", type=Path)
    args = p.parse_args()
    try:
        if args.action == "verify":
            D.require(args.package is not None, "Specify package")
            result = verify(args.package)
        else:
            D.require(args.pack is not None and args.output is not None, "Pack and new output required")
            pack = D.checked(D.load(args.pack))
            if args.action == "compile":
                result = compile_delivery(pack, args.output)
            else:
                D.require(args.workbook is not None, "Workbook required")
                result = import_profile(args.workbook, pack)
                D.write_new(args.output, result)
        print(json.dumps(result, indent=2))
        return 0
    except (ValueError, OSError, KeyError, ImportError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
